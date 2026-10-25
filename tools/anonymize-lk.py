# /// script
# requires-python = ">=3.13"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Anonymised copy of an XLSX export of the LK (section 4 of docs/specs/07-environments.md).

    uv run --script tools/anonymize-lk.py lk-export.xlsx lk-export-anon.xlsx

INN become valid synthetic ones, names "Заказчик N", keys new ones of the same shape (the
WWyy prefix of a new key is kept: it is the week of issue, not a secret); dates, quantities,
articles and serial numbers stay as they are. The same INN or name gets the same replacement in
every row, so customers with several licences stay together. Headers are found as the import
does (apps/licensing/domain/import_lk.py: case and spaces ignored, synonyms). Nothing of the
source is printed.
"""

from __future__ import annotations

import re
import secrets
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HEADERS = {
    "key": {"кодлицензии", "ключлицензии", "код", "licensekey"},
    "inn": {"инн", "inn"},
    "customer": {"заказчик", "организация", "customer"},
}
KEY = re.compile(r"[0-9A-Za-z]{4}(-[0-9A-Za-z]{4}){3}")
NEW_KEY_PREFIX = re.compile(r"\d{4}")
ALPHABET = "23456789ABCDEFGHJKLMNPQRSTUVWXYZ"
W10 = (2, 4, 10, 3, 5, 9, 4, 6, 8)
W11 = (7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
W12 = (3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
HEADER_ROWS = 10  # the header is looked for in the first rows, as in the import


def _check(digits: list[int], weights: tuple[int, ...]) -> int:
    return sum(d * w for d, w in zip(digits, weights, strict=False)) % 11 % 10


def synthetic_inn(length: int, used: set[str]) -> str:
    """A valid INN of the same length; region 99 and random digits, never one already given."""
    while True:
        body = [9, 9] + [secrets.randbelow(10) for _ in range(length - 4 if length == 12 else 7)]
        if length == 10:
            inn = body + [_check(body, W10)]
        else:
            inn = body + [_check(body, W11)]
            inn += [_check(inn, W12)]
        text = "".join(map(str, inn))
        if text not in used:
            used.add(text)
            return text


def new_key(old: str, used: set[str]) -> str:
    """A random value of the same shape: letters and digits replaced, separators kept, so a
    malformed key of the export stays malformed; the week-year prefix of a new key stays."""
    text = old.upper()
    keep = 4 if KEY.fullmatch(text) and NEW_KEY_PREFIX.fullmatch(text[:4]) else 0
    while True:
        key = text[:keep] + "".join(
            secrets.choice(ALPHABET) if c.isalnum() else c for c in text[keep:]
        )
        if key not in used and key != text:
            used.add(key)
            return key


def _normal(value: Any) -> str:
    return "".join(str(value or "").split()).lower()


def find_columns(sheet: Any) -> tuple[int, dict[str, int]]:
    for row in sheet.iter_rows(min_row=1, max_row=HEADER_ROWS):
        found = {
            name: cell.column
            for cell in row
            for name, titles in HEADERS.items()
            if _normal(cell.value) in titles
        }
        if "key" in found and "inn" in found:
            return row[0].row, found
    raise SystemExit("anonymize-lk: no header with the key and INN columns in the first rows")


def anonymize(source: Path, target: Path) -> dict[str, int]:
    book = load_workbook(source)
    sheet = book.worksheets[0]
    header_row, columns = find_columns(sheet)
    inns: dict[str, str] = {}
    names: dict[str, str] = {}
    used_inn: set[str] = set()
    used_keys: set[str] = set()
    keys = 0
    for row in sheet.iter_rows(min_row=header_row + 1):
        by_column = {cell.column: cell for cell in row}
        key = by_column.get(columns["key"])
        if key is not None and key.value not in (None, ""):
            key.value = new_key(str(key.value).strip(), used_keys)
            keys += 1
        inn = by_column.get(columns["inn"])
        if inn is not None and inn.value not in (None, ""):
            text = str(inn.value).strip()
            if text not in inns:
                inns[text] = synthetic_inn(12 if len(text) == 12 else 10, used_inn)
            inn.value = inns[text]
        if "customer" in columns:
            name = by_column.get(columns["customer"])
            if name is not None and name.value not in (None, ""):
                original = str(name.value).strip()
                names.setdefault(original, f"Заказчик {len(names) + 1}")
                name.value = names[original]
    book.save(target)
    return {"keys": keys, "inn": len(inns), "customers": len(names)}


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: anonymize-lk.py lk-export.xlsx lk-export-anon.xlsx")
    source, target = Path(sys.argv[1]), Path(sys.argv[2])
    if source.resolve() == target.resolve():
        raise SystemExit("anonymize-lk: the copy must be another file")
    counts = anonymize(source, target)
    replaced = f"{counts['keys']} keys, {counts['inn']} INN, {counts['customers']} names"
    print(f"anonymize-lk: {replaced} replaced")


if __name__ == "__main__":
    main()
