"""E16-06: spreadsheets with formulas and external links (criterion 12, section 13 of doc. 1).

An import never evaluates a formula or follows a link; an export never writes text that a
spreadsheet would run as a formula.
"""

import io
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from apps.licensing.domain import parse_lk_export
from infra.clock import FrozenClock

pytestmark = pytest.mark.django_db
HEADER = [
    "Серийный номер",
    "Код лицензии",
    "Артикул",
    "Количество",
    "ИНН",
    "Дата выдачи",
    "Заказчик",
]
EVIL = '=HYPERLINK("http://attacker.example/x","Открыть")'


def test_import_reads_no_formula_and_no_external_link(clock: FrozenClock) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(HEADER)
    sheet.append(["26-x", "2611-AAAA-BBBB-CCCC", "01", "=1+2", "7701234560", "01.11.2026", EVIL])
    # A link to another workbook: openpyxl keeps it as data and never opens it.
    sheet["H2"] = "='http://attacker.example/[book.xlsx]Sheet1'!A1"
    stream = io.BytesIO()
    book.save(stream)

    preview = parse_lk_export(stream.getvalue(), filename="evil.xlsx", now=clock.now())
    [row] = preview.rows
    # Formulas have no cached value: the cells are empty, nothing is computed or fetched.
    assert row.status == "error"
    assert not str(row.fields.get("customer", "")).startswith("=")
    assert "attacker" not in str(row.fields)


def test_export_writes_formula_text_as_text(licensed: Any, clock: FrozenClock) -> None:
    from apps.stats.domain import stats_workbook

    licensed.customer.name = EVIL
    licensed.customer.save()
    book = load_workbook(io.BytesIO(stats_workbook()))
    cells = [c for row in book.worksheets[0].iter_rows() for c in row]
    names = [c for c in cells if c.value and "attacker" in str(c.value)]
    assert names, "the customer name is in the export"
    for cell in names:
        assert cell.data_type == "s", cell.value  # a string, not a formula
