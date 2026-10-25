"""Admin API: statistics and its XLSX export (E10-01, E10-02)."""

import io
from collections.abc import Iterator
from typing import Any

import pytest
from openpyxl import load_workbook

from infra.clock import FrozenClock, use_clock
from tests.api_admin.session import sign_in

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def server_clock(clock: FrozenClock) -> Iterator[FrozenClock]:
    with use_clock(clock):
        yield clock


def test_summary_counts_and_spread(licensed: Any, make_activation: Any, clock: FrozenClock) -> None:
    variants: list[dict[str, Any]] = [
        {"product_version": "3.2.0", "os": {"name": "Astra Linux", "version": "1.8"}},
        {
            "product_version": "3.2.0",
            "os": {"name": "Windows 11", "version": "10.0"},
            "in_reserve": True,
        },
        {"product_version": "3.1.0", "status": "released_auto"},
    ]
    for n, extra in enumerate(variants):
        make_activation(
            external_id=f"act_01JABF8Q4W2N8C1R7ZB0Y5K6M{n}",
            instance_id=f"8d6a1f9e-6c39-4e57-9a3e-0e5d0f7b2c1{n}",
            **extra,
        )
    console = sign_in("operator", clock)

    body = console.get("/api/admin/stats", {"customer_id": str(licensed.customer_id)}).json()
    assert (body["licenses"], body["limit"], body["active"], body["in_reserve"], body["free"]) == (
        1,
        10,
        2,
        1,
        8,
    )
    assert body["released_auto"] == 1
    assert body["by_version"] == {"3.2.0": 2}
    assert body["by_os"] == {"Astra Linux 1.8": 1, "Windows 11 10.0": 1}
    [row] = body["rows"]
    assert (row["product"], row["customer_name"]) == ("alteroffice-desktop", "ООО «Пример»")
    assert row["last_seen_at"] is not None


def test_xlsx_export_has_summary_and_activations(
    licensed: Any, make_activation: Any, clock: FrozenClock
) -> None:
    make_activation()
    console = sign_in("operator", clock)
    response = console.get("/api/admin/stats/export")
    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="almi-stats.xlsx"'
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert book.sheetnames == ["Сводка", "Активации"]
    summary = list(book["Сводка"].iter_rows(values_only=True))
    assert summary[1][:2] == ("ООО «Пример»", "alteroffice-desktop")
    assert summary[-1][0] == "Итого"
    assert len(list(book["Активации"].iter_rows(values_only=True))) == 2


def test_viewer_sees_only_its_statistics(clock: FrozenClock, licensed: Any) -> None:
    stats = sign_in("customer_viewer", clock).get("/api/admin/stats").json()
    assert stats["licenses"] == 0  # licensed belongs to another customer (E07-08)
