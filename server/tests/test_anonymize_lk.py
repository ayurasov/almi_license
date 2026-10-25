"""tools/anonymize-lk.py: an anonymised copy of the LK export keeps everything but the secrets."""

import importlib.util
import re
from pathlib import Path
from types import ModuleType
from typing import Any

from openpyxl import load_workbook

from apps.licensing.domain import inn_is_valid, normalize_key

TOOL = Path(__file__).resolve().parents[2] / "tools" / "anonymize-lk.py"
KEY = re.compile(r"[0-9A-Z]{4}(-[0-9A-Z]{4}){3}")
SAMPLE = Path(__file__).resolve().parent / "fixtures" / "lk-export-sample.xlsx"


def tool() -> ModuleType:
    spec = importlib.util.spec_from_file_location("anonymize_lk", TOOL)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    sheet = load_workbook(path).worksheets[0]
    values = [list(r) for r in sheet.iter_rows(values_only=True)]
    header = next(
        i for i, r in enumerate(values) if any(str(v).strip().lower() == "инн" for v in r if v)
    )
    return [str(v or "") for v in values[header]], values[header + 1 :]


def test_secrets_are_replaced_and_the_rest_is_kept(tmp_path: Path) -> None:
    target = tmp_path / "anon.xlsx"
    counts = tool().anonymize(SAMPLE, target)
    header, before = rows(SAMPLE)
    _, after = rows(target)
    lowered = [h.strip().lower() for h in header]
    key, inn = lowered.index("код лицензии"), lowered.index("инн")
    assert counts["keys"] > 0 and len(before) == len(after)
    replacements: dict[object, object] = {}
    for old, new in zip(before, after, strict=True):
        for column, (a, b) in enumerate(zip(old, new, strict=True)):
            if column not in (key, inn) and lowered[column] not in ("заказчик", "организация"):
                assert a == b  # dates, quantities, articles, serial numbers
        if old[key]:
            assert new[key] != old[key] and len(str(new[key])) == len(str(old[key]).strip())
            if KEY.fullmatch(str(old[key]).strip().upper()):
                assert normalize_key(str(new[key]))
        if old[inn]:
            assert new[inn] != old[inn] and inn_is_valid(str(new[inn]))
            # The same customer keeps one replacement in all its rows.
            assert replacements.setdefault(old[inn], new[inn]) == new[inn]
