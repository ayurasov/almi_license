"""Statistics of licences for the console (E10-01) and its XLSX export (E10-02).

Per licence: limit, active, in reserve, free, released automatically, conflicts, last contact,
spread by product version and OS; totals over the filter. The export has a summary sheet and an
activations sheet (at most 100 000 rows), written in streaming mode.
"""

import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any
from uuid import UUID

from openpyxl import Workbook

from apps.catalog.domain import product_info
from apps.licensing.domain import counts_active_only, find_licenses
from apps.stats.domain.usage import latest_usage

MAX_ACTIVATION_ROWS = 100_000


@dataclass(frozen=True)
class LicenseStats:
    license_id: UUID
    customer_id: UUID
    customer_name: str
    product: str
    type: str
    status: str
    limit: int
    active: int
    in_reserve: int
    free: int
    released_auto: int
    conflicts: int
    valid_until: datetime | None
    last_seen_at: datetime | None
    by_version: dict[str, int]
    by_os: dict[str, int]
    users: int | None = None
    user_limit: int | None = None
    users_exceeded: bool = False
    state: str = ""
    local_active: int = 0


@dataclass(frozen=True)
class Summary:
    licenses: int
    limit: int
    active: int
    in_reserve: int
    free: int
    released_auto: int
    conflicts: int
    last_seen_at: datetime | None
    by_version: dict[str, int]
    by_os: dict[str, int]
    rows: list[LicenseStats]
    by_state: dict[str, int] = field(default_factory=dict)


def license_stats(*, customer_id: UUID | None = None, product_id: UUID | None = None) -> Summary:
    # activation records usage through stats
    from apps.activation.domain import active_activation_ids, usage_by_license
    from apps.packages.domain import local_active

    licenses = find_licenses(customer_id=customer_id, product_id=product_id)
    usage = usage_by_license([lic.id for lic in licenses])
    active_ids = active_activation_ids([lic.id for lic in licenses])
    counters = latest_usage([i for ids in active_ids.values() for i in ids])
    on_local = local_active([lic.id for lic in licenses])  # reports of local servers (E09-06)
    products: dict[UUID, str] = {}
    rows: list[LicenseStats] = []
    versions: Counter[str] = Counter()
    systems: Counter[str] = Counter()
    for lic in licenses:
        used = usage[lic.id]
        users, user_limit = _users(
            lic.limits, [counters[i] for i in active_ids[lic.id] if i in counters]
        )
        if lic.product_id not in products:
            products[lic.product_id] = product_info(lic.product_id).code
        rows.append(
            LicenseStats(
                license_id=lic.id,
                customer_id=lic.customer_id,
                customer_name=lic.customer_name,
                product=products[lic.product_id],
                type=lic.type,
                status=lic.status,
                limit=lic.limit,
                active=used.active,
                in_reserve=used.in_reserve,
                free=max(lic.limit - used.active, 0),
                released_auto=used.released_auto,
                conflicts=used.conflicts,
                valid_until=lic.valid_until,
                last_seen_at=used.last_seen_at,
                by_version=used.by_version,
                by_os=used.by_os,
                users=users,
                user_limit=user_limit,
                users_exceeded=users is not None and user_limit is not None and users > user_limit,
                state=lic.state,
                local_active=on_local.get(lic.id, 0),
            )
        )
        versions.update(used.by_version)
        systems.update(used.by_os)
    seen = [r.last_seen_at for r in rows if r.last_seen_at is not None]
    return Summary(
        licenses=len(rows),
        limit=sum(r.limit for r in rows),
        active=sum(r.active for r in rows),
        in_reserve=sum(r.in_reserve for r in rows),
        free=sum(r.free for r in rows),
        released_auto=sum(r.released_auto for r in rows),
        conflicts=sum(r.conflicts for r in rows),
        last_seen_at=max(seen) if seen else None,
        by_version=dict(versions),
        by_os=dict(systems),
        rows=rows,
        by_state=dict(Counter(r.state for r in rows if r.state)),
    )


def _users(limits: dict[str, Any], counters: list[dict[str, Any]]) -> tuple[int | None, int | None]:
    """Users reported by the instances of a server product and the limit they are counted against.

    only_count_active (or a limit of active users alone) counts active users (E11-02).
    """
    field = "active_users" if counts_active_only(limits) else "named_users"
    reported = [c[field] for c in counters if isinstance(c.get(field), int)]
    limit = limits.get("named_users", limits.get("active_users"))
    return (sum(reported) if reported else None), limit


def _cell(value: Any) -> Any:
    # Excel has no time zones: dates are written in UTC without the offset.
    return value.replace(tzinfo=None) if isinstance(value, datetime) else value


def stats_workbook(*, customer_id: UUID | None = None, product_id: UUID | None = None) -> bytes:
    from apps.activation.domain import activations_of  # activation records usage through stats

    summary = license_stats(customer_id=customer_id, product_id=product_id)
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Сводка")
    sheet.append(
        [
            "Заказчик",
            "Продукт",
            "Лицензия",
            "Тип",
            "Статус",
            "Мест",
            "Активно",
            "В резерве",
            "Свободно",
            "Освобождено автоматически",
            "Конфликты",
            "Действует до (UTC)",
            "Последний контакт (UTC)",
            "Пользователей",
            "Лимит пользователей",
        ]
    )
    for row in summary.rows:
        sheet.append(
            [
                _cell(v)
                for v in (
                    row.customer_name,
                    row.product,
                    str(row.license_id),
                    row.type,
                    row.status,
                    row.limit,
                    row.active,
                    row.in_reserve,
                    row.free,
                    row.released_auto,
                    row.conflicts,
                    row.valid_until,
                    row.last_seen_at,
                    row.users,
                    row.user_limit,
                )
            ]
        )
    sheet.append([])
    sheet.append(
        [
            "Итого",
            "",
            f"{summary.licenses} лицензий",
            "",
            "",
            summary.limit,
            summary.active,
            summary.in_reserve,
            summary.free,
            summary.released_auto,
            summary.conflicts,
            "",
            _cell(summary.last_seen_at),
        ]
    )
    activations = book.create_sheet("Активации")
    activations.append(
        [
            "Лицензия",
            "Активация",
            "Экземпляр",
            "Имя",
            "Статус",
            "В резерве",
            "Версия продукта",
            "ОС",
            "Активирована (UTC)",
            "Последний контакт (UTC)",
            "Освобождена (UTC)",
        ]
    )
    for item in activations_of([r.license_id for r in summary.rows], limit=MAX_ACTIVATION_ROWS):
        os = f"{item.os.get('name', '')} {item.os.get('version', '')}".strip()
        activations.append(
            [
                _cell(v)
                for v in (
                    str(item.license_id),
                    item.activation_id,
                    str(item.instance_id),
                    item.instance_name,
                    item.status,
                    "да" if item.in_reserve else "",
                    item.product_version,
                    os,
                    item.activated_at,
                    item.last_seen_at,
                    item.released_at,
                )
            ]
        )
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
