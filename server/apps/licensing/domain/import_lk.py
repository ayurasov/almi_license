"""Preview of an XLSX export of the LK (E06-01, E06-02).

Columns — section 14.1 of document 2, articles — section 8 of document 4 (decisions of
23.10.2026: the article is dash-separated segments AA-BB-CC-D-EE-FF-…, AA first, FF sixth;
headers are matched without case and spaces, with synonyms). Nothing is written to licenses:
the preview is stored as a draft ImportSession for 24 hours. Keys appear only as mask and hash.
"""

import calendar
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta, timezone
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.catalog.domain import edition_code, metric_code, product_info
from apps.licensing.domain.keys import InvalidKeyFormat, key_hash, key_mask, normalize_key
from apps.licensing.models import ArticleMapping, ImportSession, LicenseKey

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
SESSION_TTL = timedelta(hours=24)
# Dates of the LK are Moscow dates; Moscow time has been UTC+3 without DST since 2014.
MSK = timezone(timedelta(hours=3))
AA_SEGMENT, FF_SEGMENT, MIN_SEGMENTS = 0, 5, 6
TERMS_MONTHS = {"01": 1, "03": 3, "06": 6, "12": 12, "24": 24, "36": 36}
PERPETUAL_TERM = "99"

# field -> (title for messages, accepted headers normalised: lower case, no whitespace)
COLUMNS: dict[str, tuple[str, frozenset[str]]] = {
    "serial": ("Серийный номер", frozenset({"серийныйномер", "серийный№", "serial"})),
    "key": ("Код лицензии", frozenset({"кодлицензии", "ключлицензии", "код", "licensekey"})),
    "article": ("Артикул", frozenset({"артикул", "article"})),
    "quantity": ("Количество", frozenset({"количество", "кол-во", "totallicenses", "total"})),
    "inn": ("ИНН", frozenset({"инн", "inn"})),
    "issued": ("Дата выдачи", frozenset({"датавыдачи", "датавыпуска", "issuedate"})),
    "customer": ("Заказчик", frozenset({"заказчик", "организация", "customer"})),
}
OPTIONAL = frozenset({"customer"})
_SERIAL = re.compile(r"(\d{2})-(.+)-(\d{3})-(\d{4})")
_SEGMENT = re.compile(r"[0-9A-Za-z]+")
_INN = re.compile(r"[0-9]{10}([0-9]{2})?")
_TEXT_DATE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


class ImportFileError(ValueError):
    """The whole file is refused; the message is for the operator (Russian)."""


@dataclass(frozen=True)
class ImportRow:
    row: int  # spreadsheet row number
    status: str  # new | update | skip | error
    errors: list[str]
    fields: dict[str, Any]
    key_hash: str | None = field(default=None, repr=False)


@dataclass(frozen=True)
class ImportPreview:
    session_id: UUID
    filename: str
    rows: list[ImportRow]
    summary: dict[str, int]


@dataclass(frozen=True)
class _Mapping:
    product_id: UUID
    edition_id: UUID | None
    metric_id: UUID
    product: str
    edition: str | None
    metric: str


def _header_name(value: object) -> str:
    return "".join(str(value or "").split()).lower()


def _columns(header: tuple[object, ...]) -> dict[str, int]:
    names = [_header_name(value) for value in header]
    found: dict[str, int] = {}
    for name, (title, accepted) in COLUMNS.items():
        index = next((i for i, value in enumerate(names) if value in accepted), None)
        if index is None and name not in OPTIONAL:
            raise ImportFileError(f"Нет обязательной колонки: {title}")
        if index is not None:
            found[name] = index
    return found


def _read(data: bytes) -> tuple[tuple[object, ...], list[tuple[int, tuple[object, ...]]]]:
    if len(data) > MAX_BYTES:
        raise ImportFileError("Файл больше 20 МБ")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        raise ImportFileError("Файл не является книгой XLSX") from None
    try:
        rows = book.worksheets[0].iter_rows(values_only=True) if book.worksheets else iter(())
        header: tuple[object, ...] | None = None
        body: list[tuple[int, tuple[object, ...]]] = []
        for number, values in enumerate(rows, start=1):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            if header is None:
                header = values
                continue
            if len(body) == MAX_ROWS:
                raise ImportFileError(f"В файле больше {MAX_ROWS} строк")
            body.append((number, values))
    finally:
        book.close()
    if header is None or not body:
        raise ImportFileError("В файле нет данных")
    return header, body


def _add_months(day: date, months: int) -> date:
    month = day.month - 1 + months
    year, month = day.year + month // 12, month % 12 + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def _issue_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = _TEXT_DATE.fullmatch(str(value or "").strip())
    if match is None:
        return None
    try:
        return date(int(match[3]), int(match[2]), int(match[1]))
    except ValueError:
        return None


def _quantity(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value.strip())
    return value if isinstance(value, int) and value > 0 else None


def _rfc3339(moment: datetime | None) -> str | None:
    return None if moment is None else moment.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _mappings() -> dict[str, _Mapping]:
    result = {}
    for row in ArticleMapping.objects.all():
        result[row.aa] = _Mapping(
            product_id=row.product_id,
            edition_id=row.edition_id,
            metric_id=row.metric_id,
            product=product_info(row.product_id).code,
            edition=edition_code(row.edition_id),
            metric=metric_code(row.metric_id),
        )
    return result


def _cell(values: tuple[object, ...], columns: dict[str, int], name: str) -> object:
    index = columns.get(name)
    return values[index] if index is not None and index < len(values) else None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_row(
    values: tuple[object, ...], columns: dict[str, int], mappings: dict[str, _Mapping]
) -> tuple[list[str], dict[str, Any], str | None, _Mapping | None]:
    errors: list[str] = []
    fields: dict[str, Any] = {}
    serial = _text(_cell(values, columns, "serial"))
    article = _text(_cell(values, columns, "article"))
    fields["serial"] = serial
    serial_match = _SERIAL.fullmatch(serial)
    if serial_match is None:
        errors.append("Неверный серийный номер")

    digest = None
    try:
        raw_key = _text(_cell(values, columns, "key"))
        digest = key_hash(raw_key)
        fields["key_mask"] = key_mask(normalize_key(raw_key))
    except InvalidKeyFormat:
        errors.append("Неверный формат кода лицензии")

    fields["article"] = article
    segments = article.split("-")
    mapping = None
    if len(segments) < MIN_SEGMENTS or not all(_SEGMENT.fullmatch(s) for s in segments):
        errors.append("Неверный артикул")
    else:
        aa, ff = segments[AA_SEGMENT], segments[FF_SEGMENT]
        fields |= {"aa": aa, "ff": ff}
        if serial_match is not None and serial_match[2] != article:
            errors.append("Серийный номер не соответствует артикулу")
        mapping = mappings.get(aa)
        if mapping is None:
            errors.append(f"Артикул не сопоставлен: {aa}")
        else:
            fields |= {
                "product": mapping.product,
                "edition": mapping.edition,
                "metric": mapping.metric,
            }
        if ff != PERPETUAL_TERM and ff not in TERMS_MONTHS:
            errors.append(f"Неизвестный срок FF: {ff}")

    quantity = _quantity(_cell(values, columns, "quantity"))
    if quantity is None:
        errors.append("Количество должно быть целым числом больше нуля")
    inn = _text(_cell(values, columns, "inn"))
    if _INN.fullmatch(inn) is None:
        errors.append("Неверный ИНН")
    issued = _issue_date(_cell(values, columns, "issued"))
    if issued is None:
        errors.append("Неверная дата выдачи")

    term = fields.get("ff")
    if issued is not None and (term == PERPETUAL_TERM or term in TERMS_MONTHS):
        valid_from = datetime.combine(issued, time(), MSK)
        valid_until = (
            None
            if term == PERPETUAL_TERM
            else datetime.combine(
                _add_months(issued, TERMS_MONTHS[str(term)]) - timedelta(days=1),
                time(23, 59, 59),
                MSK,
            )
        )
        fields |= {
            "type": "perpetual" if term == PERPETUAL_TERM else "subscription",
            "valid_from": _rfc3339(valid_from),
            "valid_until": _rfc3339(valid_until),
        }
    fields |= {
        "limit": quantity,
        "inn": inn,
        "customer": _text(_cell(values, columns, "customer")),
    }
    return errors, fields, digest, mapping


def _status(
    fields: dict[str, Any], mapping: _Mapping, existing: LicenseKey | None
) -> tuple[str, list[str]]:
    if existing is None:
        return "new", []
    license_ = existing.license
    if license_.customer.inn != fields["inn"]:
        return "error", ["Ключ принадлежит другому заказчику"]
    same = (
        license_.product_id == mapping.product_id
        and license_.edition_id == mapping.edition_id
        and license_.type == fields["type"]
        and license_.limit == fields["limit"]
        and _rfc3339(license_.valid_from) == fields["valid_from"]
        and _rfc3339(license_.valid_until) == fields["valid_until"]
    )
    return ("skip" if same else "update"), []


def parse_lk_export(data: bytes, *, filename: str, now: datetime) -> ImportPreview:
    """Parse the export and store a draft ImportSession; ImportFileError refuses the file."""
    header, body = _read(data)
    columns = _columns(header)
    mappings = _mappings()
    parsed = [(number, *_parse_row(values, columns, mappings)) for number, values in body]

    digests = [digest for _, errors, _, digest, _ in parsed if digest is not None]
    existing = {
        key.key_hash: key
        for key in LicenseKey.objects.filter(key_hash__in=digests).select_related(
            "license__customer"
        )
    }
    first_row: dict[str, int] = {}
    rows: list[ImportRow] = []
    for number, errors, fields, digest, mapping in parsed:
        if digest is not None:
            if digest in first_row:
                errors.append(f"Код лицензии повторяется в файле: строка {first_row[digest]}")
            else:
                first_row[digest] = number
        if errors or mapping is None:
            rows.append(ImportRow(number, "error", errors, fields, digest))
            continue
        status, reasons = _status(fields, mapping, existing.get(str(digest)))
        rows.append(ImportRow(number, status, reasons, fields, digest))

    summary = {
        kind: sum(row.status == kind for row in rows) for kind in ("new", "update", "skip", "error")
    }
    session = ImportSession.objects.create(
        filename=filename,
        rows=[
            {
                "row": r.row,
                "status": r.status,
                "errors": r.errors,
                "fields": r.fields,
                "key_hash": r.key_hash,
            }
            for r in rows
        ],
        summary=summary,
        created_at=now,
        expires_at=now + SESSION_TTL,
    )
    return ImportPreview(session_id=session.id, filename=filename, rows=rows, summary=summary)
